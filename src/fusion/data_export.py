"""
File: src/fusion/data_export.py
Purpose: Generates downloadable diagnosis reports in multiple formats (.txt, .csv, .pdf, .docx).
Why we need it: After a diagnosis, users should be able to save and share their results
                in whatever format suits them — plain text for simplicity, CSV for data logging,
                PDF for a formal-looking report, and DOCX for easy editing.
"""

import io
import csv
from datetime import datetime

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
)
from reportlab.lib.enums import TA_CENTER, TA_LEFT

import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ==============================================================================
# SHARED HELPERS
# ==============================================================================

def _format_datetime() -> str:
    """Returns a formatted timestamp string for report headers."""
    return datetime.now().strftime("%B %d, %Y at %I:%M %p")


def _build_report_content(result: dict, symptom_text: str) -> dict:
    """
    Extracts and normalises all needed data from a fusion result dict into a 
    flat structure ready for any format renderer to consume.
    Why we need it: Avoids duplicating parsing logic in every export function.
    """
    from src.fusion.narrative import generate_specialist_narrative
    narrative = generate_specialist_narrative(result, symptom_text)

    clean_name = result.get("prediction", "Unknown").replace("_", " ").title()
    confidence = result.get("confidence", 0.0) * 100
    status = result.get("status", "unknown")
    probs = result.get("probabilities", {})

    # Sort probabilities descending for display
    sorted_probs = sorted(probs.items(), key=lambda x: x[1], reverse=True)
    formatted_probs = [(k.replace("_", " ").title(), f"{v * 100:.2f}%") for k, v in sorted_probs]

    return {
        "prediction": clean_name,
        "formal_name": narrative["formal_name"],
        "category": narrative["category"],
        "urgency": narrative["urgency"],
        "certainty_label": narrative["certainty_label"],
        "confidence": f"{confidence:.2f}%",
        "status": status,
        "symptom_text": symptom_text.strip() if symptom_text else "Not provided",
        "timestamp": _format_datetime(),
        "probabilities": formatted_probs,
        "lead_paragraph": narrative["lead_paragraph"],
        "visual_findings": narrative["visual_findings"],
        "differential_analysis": narrative["differential_analysis"],
        "clinical_actions": narrative["clinical_actions"],
        "red_flags": narrative.get("red_flags", []),
    }


# ==============================================================================
# TXT EXPORT
# ==============================================================================

def export_txt(result: dict, symptom_text: str) -> bytes:
    """
    Generates a plain-text diagnosis report.
    Why we need it: The simplest, most universally readable format — no dependencies,
                    works on any device or system.
    Returns raw bytes for Streamlit's st.download_button.
    """
    data = _build_report_content(result, symptom_text)
    lines = [
        "=" * 70,
        "  MULTIMODAL SKIN DISEASE DIAGNOSIS REPORT (SPECIALIST CONSULTATION)",
        "=" * 70,
        f"  Generated: {data['timestamp']}",
        "",
        "DISCLAIMER: This report is for academic/demonstrative purposes",
        "only. It does not replace professional medical consultation.",
        "",
        "-" * 70,
        "PATIENT INPUT",
        "-" * 70,
        f"Symptoms: {data['symptom_text']}",
        "",
        "-" * 70,
        "DIAGNOSTIC SUMMARY",
        "-" * 70,
        f"Primary Indication : {data['formal_name']}",
        f"Taxonomy Category  : {data['category']}",
        f"Diagnostic Certainty: {data['confidence']} ({data['certainty_label']})",
        f"Triage Advisory    : {data['urgency']}",
        f"System Status      : {data['status'].replace('_', ' ').title()}",
        "",
        "-" * 70,
        "SPECIALIST CONSULTATION ASSESSMENT",
        "-" * 70,
        data['lead_paragraph'].replace("**", ""),
        "",
        "OBSERVED CLINICAL & MORPHOLOGICAL INDICATORS:",
    ]
    for finding in data["visual_findings"]:
        lines.append(f"  • {finding}")

    lines.append("")
    lines.append("DIFFERENTIAL DIAGNOSIS CONSIDERATIONS:")
    lines.append(f"  {data['differential_analysis']}")

    lines.append("")
    lines.append("RECOMMENDED CLINICAL NEXT STEPS:")
    for i, action in enumerate(data["clinical_actions"], 1):
        lines.append(f"  {i}. {action}")

    if data.get("red_flags"):
        lines.append("")
        lines.append("RED-FLAG SYMPTOMS REQUIRING IMMEDIATE EMERGENCY CARE:")
        for rf in data["red_flags"]:
            lines.append(f"  🚨 {rf}")

    lines.extend([
        "",
        "-" * 70,
        "CLASS PROBABILITIES (Highest to Lowest)",
        "-" * 70,
    ])
    for condition, prob in data["probabilities"]:
        lines.append(f"  {condition:<45} {prob}")
    lines += [
        "",
        "=" * 70,
        "  Please consult a qualified dermatologist for a confirmed diagnosis.",
        "=" * 70,
    ]
    return "\n".join(lines).encode("utf-8")


# ==============================================================================
# CSV EXPORT
# ==============================================================================

def export_csv(result: dict, symptom_text: str) -> bytes:
    """
    Generates a CSV report with a metadata section and a probability table.
    Why we need it: Structured data format — useful for logging multiple diagnoses
                    or importing results into a spreadsheet.
    Returns raw bytes for Streamlit's st.download_button.
    """
    data = _build_report_content(result, symptom_text)
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    # Metadata block
    writer.writerow(["Field", "Value"])
    writer.writerow(["Generated", data["timestamp"]])
    writer.writerow(["Symptoms", data["symptom_text"]])
    writer.writerow(["Predicted Condition", data["prediction"]])
    writer.writerow(["Primary Indication", data["formal_name"]])
    writer.writerow(["Taxonomy Category", data["category"]])
    writer.writerow(["Diagnostic Certainty", data["confidence"]])
    writer.writerow(["Certainty Level", data["certainty_label"]])
    writer.writerow(["Triage Advisory", data["urgency"]])
    writer.writerow(["Specialist Assessment", data["lead_paragraph"].replace("**", "")])
    writer.writerow(["Differential Notes", data["differential_analysis"]])
    writer.writerow(["Recommended Actions", "; ".join(data["clinical_actions"])])
    writer.writerow(["Status", data["status"].replace("_", " ").title()])
    writer.writerow([])

    # Probability table
    writer.writerow(["Condition", "Probability"])
    for condition, prob in data["probabilities"]:
        writer.writerow([condition, prob])

    writer.writerow([])
    writer.writerow(["Disclaimer", "For academic/demonstrative purposes only. Not medical advice."])

    return buffer.getvalue().encode("utf-8")


# ==============================================================================
# PDF EXPORT
# ==============================================================================

def export_pdf(result: dict, symptom_text: str) -> bytes:
    """
    Generates a clean, formatted A4 PDF report using ReportLab.
    Why we need it: PDF is the standard for shareable, print-ready formal documents
                    — this is the most likely format a user would hand to a doctor.
    Returns raw bytes for Streamlit's st.download_button.
    """
    data = _build_report_content(result, symptom_text)
    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
    )

    styles = getSampleStyleSheet()
    # Custom styles
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        fontSize=18,
        textColor=colors.HexColor("#1E88E5"),
        alignment=TA_CENTER,
        spaceAfter=4,
    )
    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        parent=styles["Normal"],
        fontSize=10,
        textColor=colors.grey,
        alignment=TA_CENTER,
        spaceAfter=2,
    )
    section_heading_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontSize=12,
        textColor=colors.HexColor("#1565C0"),
        spaceBefore=10,
        spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        spaceAfter=4,
    )
    disclaimer_style = ParagraphStyle(
        "Disclaimer",
        parent=styles["Normal"],
        fontSize=8,
        textColor=colors.grey,
        alignment=TA_CENTER,
        spaceBefore=12,
    )

    story = []

    # --- Header ---
    story.append(Paragraph("🩺 Skin Disease Diagnosis Report", title_style))
    story.append(Paragraph(f"Generated: {data['timestamp']}", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#1E88E5"), spaceAfter=10))

    # --- Patient Input ---
    story.append(Paragraph("Patient Symptom Description", section_heading_style))
    story.append(Paragraph(data["symptom_text"], body_style))
    story.append(Spacer(1, 6))

    # --- Diagnosis Result ---
    story.append(Paragraph("Diagnostic Summary", section_heading_style))
    result_table_data = [
        ["Field", "Value"],
        ["Primary Indication", data["formal_name"]],
        ["Taxonomy Category", data["category"]],
        ["Diagnostic Certainty", f"{data['confidence']} ({data['certainty_label']})"],
        ["Triage Advisory", data["urgency"]],
        ["System Status", data["status"].replace("_", " ").title()],
    ]
    result_table = Table(result_table_data, colWidths=[55 * mm, 115 * mm])
    result_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E88E5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#F5F5F5"), colors.white]),
        ("FONTSIZE", (0, 1), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#BDBDBD")),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(result_table)
    story.append(Spacer(1, 8))

    # --- Specialist Consultation Assessment ---
    story.append(Paragraph("Specialist Consultation Assessment", section_heading_style))
    story.append(Paragraph(data["lead_paragraph"].replace("**", ""), body_style))
    story.append(Spacer(1, 4))

    # Observed Indicators & Recommended Actions
    story.append(Paragraph("<b>Observed Morphological Indicators:</b>", body_style))
    for finding in data["visual_findings"]:
        story.append(Paragraph(f"• {finding}", body_style))
    story.append(Spacer(1, 4))

    story.append(Paragraph("<b>Recommended Clinical Next Steps:</b>", body_style))
    for i, action in enumerate(data["clinical_actions"], 1):
        story.append(Paragraph(f"{i}. {action}", body_style))
    story.append(Spacer(1, 8))

    # --- Class Probabilities ---
    story.append(Paragraph("Class Probabilities", section_heading_style))
    prob_table_data = [["Condition", "Probability"]] + list(data["probabilities"])
    prob_table = Table(prob_table_data, colWidths=[120 * mm, 50 * mm])
    prob_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1565C0")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#E3F2FD"), colors.white]),
        ("FONTSIZE", (0, 1), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#BDBDBD")),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(prob_table)

    # --- Disclaimer ---
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.grey, spaceBefore=12))
    story.append(Paragraph(
        "⚠️ DISCLAIMER: This report is generated by an academic prototype and is for "
        "demonstrative purposes only. It does NOT constitute medical advice. "
        "Please consult a qualified dermatologist for a confirmed diagnosis.",
        disclaimer_style
    ))

    doc.build(story)
    buffer.seek(0)
    return buffer.read()


# ==============================================================================
# DOCX EXPORT
# ==============================================================================

def export_docx(result: dict, symptom_text: str) -> bytes:
    """
    Generates a formatted Word document (.docx) diagnosis report.
    Why we need it: Word documents are the most editable format — a doctor or
                    patient could easily annotate the report before filing it.
    Returns raw bytes for Streamlit's st.download_button.
    """
    data = _build_report_content(result, symptom_text)
    doc = docx.Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # --- Title ---
    title = doc.add_heading("Skin Disease Diagnosis Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.color.rgb = RGBColor(0x1E, 0x88, 0xE5)

    subtitle = doc.add_paragraph(f"Generated: {data['timestamp']}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(10)
    subtitle.runs[0].font.color.rgb = RGBColor(0x75, 0x75, 0x75)

    doc.add_paragraph()  # Spacer

    # --- Patient Input ---
    doc.add_heading("Patient Symptom Description", level=2)
    p = doc.add_paragraph(data["symptom_text"])
    p.runs[0].font.size = Pt(10)

    # --- Diagnosis Result ---
    doc.add_heading("Diagnostic Summary", level=2)
    headers = [
        ("Primary Indication", data["formal_name"]),
        ("Taxonomy Category", data["category"]),
        ("Diagnostic Certainty", f"{data['confidence']} ({data['certainty_label']})"),
        ("Triage Advisory", data["urgency"]),
        ("System Status", data["status"].replace("_", " ").title())
    ]
    result_table = doc.add_table(rows=len(headers) + 1, cols=2)
    result_table.style = "Table Grid"

    # Header row
    hdr_cells = result_table.rows[0].cells
    hdr_cells[0].text = "Field"
    hdr_cells[1].text = "Value"
    for cell in hdr_cells:
        run = cell.paragraphs[0].runs[0]
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell._tc.get_or_add_tcPr()

    for i, (field, value) in enumerate(headers):
        row_cells = result_table.rows[i + 1].cells
        row_cells[0].text = field
        row_cells[1].text = value

    doc.add_paragraph()  # Spacer

    # --- Specialist Consultation Assessment ---
    doc.add_heading("Specialist Consultation Assessment", level=2)
    p_lead = doc.add_paragraph(data["lead_paragraph"].replace("**", ""))
    p_lead.runs[0].font.size = Pt(10)

    doc.add_heading("Observed Morphological Indicators", level=3)
    for finding in data["visual_findings"]:
        p_f = doc.add_paragraph(f"• {finding}")
        p_f.runs[0].font.size = Pt(10)

    doc.add_heading("Recommended Clinical Next Steps", level=3)
    for i, action in enumerate(data["clinical_actions"], 1):
        p_a = doc.add_paragraph(f"{i}. {action}")
        p_a.runs[0].font.size = Pt(10)

    doc.add_paragraph()  # Spacer

    # --- Class Probabilities ---
    doc.add_heading("Class Probabilities", level=2)
    prob_table = doc.add_table(rows=len(data["probabilities"]) + 1, cols=2)
    prob_table.style = "Table Grid"

    # Header row
    hdr_cells = prob_table.rows[0].cells
    hdr_cells[0].text = "Condition"
    hdr_cells[1].text = "Probability"
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].bold = True

    for i, (condition, prob) in enumerate(data["probabilities"]):
        row_cells = prob_table.rows[i + 1].cells
        row_cells[0].text = condition
        row_cells[1].text = prob

    # --- Disclaimer ---
    doc.add_paragraph()
    disclaimer = doc.add_paragraph(
        "⚠️ DISCLAIMER: This report is generated by an academic prototype and is for "
        "demonstrative purposes only. It does NOT constitute medical advice. "
        "Please consult a qualified dermatologist for a confirmed diagnosis."
    )
    disclaimer.runs[0].font.size = Pt(8)
    disclaimer.runs[0].font.color.rgb = RGBColor(0x75, 0x75, 0x75)
    disclaimer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()
